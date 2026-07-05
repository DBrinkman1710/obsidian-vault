from __future__ import annotations

import csv
import io
import json
import uuid
from typing import Annotated, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth.dependencies import AdminUser, CurrentUser
from app.database import get_db
from app.modules.contacts import service
from app.modules.contacts.models import Contact
from app.modules.contacts.schemas import (
    CallAnalyzeRequest,
    CallAnalyzeResponse,
    CallLogSaveRequest,
    CallLogSaveResponse,
    CompanyContactOut,
    CompanyCreate,
    CompanyOut,
    CompanyUpdate,
    ContactCreate,
    ContactLabelCreate,
    ContactLabelOut,
    ContactLabelUpdate,
    ContactList,
    ContactOut,
    ContactUpdate,
    ImportPreview,
    ImportResult,
)

# Yippie contact fields a file column can be mapped onto. full_name is required.
_IMPORT_TARGET_FIELDS = {"full_name", "email", "phone", "company", "notes"}

router = APIRouter(prefix="/contacts", tags=["contacts"])

DB = Annotated[AsyncSession, Depends(get_db)]


def _parse_import_file(filename: str, content: bytes) -> list[dict]:
    """Parse an uploaded CSV / JSON / XLSX file into a list of row dicts."""
    name = (filename or "").lower()

    if name.endswith(".json"):
        try:
            data = json.loads(content.decode("utf-8-sig"))
        except (json.JSONDecodeError, UnicodeDecodeError) as e:
            raise HTTPException(status_code=400, detail=f"Invalid JSON file: {e}")
        if isinstance(data, dict):
            data = data.get("contacts", data.get("items", []))
        if not isinstance(data, list):
            raise HTTPException(status_code=400, detail="JSON must be a list of contact objects")
        return [r for r in data if isinstance(r, dict)]

    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise HTTPException(status_code=400, detail="XLSX support is not available on the server")
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Invalid XLSX file: {e}")
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            return []
        keys = [str(h).strip().lower() if h is not None else "" for h in header]
        out: list[dict] = []
        for raw in rows_iter:
            if raw is None or all(v is None for v in raw):
                continue
            # Use min() to guard against short rows (read_only mode omits trailing empty cells)
            out.append({
                keys[i]: (str(raw[i]) if raw[i] is not None else None)
                for i in range(min(len(keys), len(raw)))
                if keys[i]
            })
        return out

    # default: CSV
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise HTTPException(status_code=400, detail="CSV file must be UTF-8 encoded")
    reader = csv.DictReader(io.StringIO(text))
    return [{(k or "").strip().lower(): v for k, v in row.items()} for row in reader]


@router.get("", response_model=ContactList)
async def list_contacts(
    current_user: CurrentUser,
    db: DB,
    search: Optional[str] = Query(None),
    skip: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=200),
    label_id: Optional[uuid.UUID] = Query(None),
    company_id: Optional[uuid.UUID] = Query(None),
    include_deleted: bool = Query(False),
):
    items, total = await service.list_contacts(
        db, current_user.tenant_id, search, skip, limit, label_id, company_id, include_deleted
    )
    return ContactList(items=items, total=total)


@router.post("", response_model=ContactOut, status_code=status.HTTP_201_CREATED)
async def create_contact(body: ContactCreate, current_user: CurrentUser, db: DB):
    return await service.create_contact(db, current_user.tenant_id, current_user.id, body)


# Label, company, import and export routes MUST stay above the dynamic /{contact_id}
# routes — FastAPI matches in declaration order and "labels"/"companies"/"import"/
# "export" would otherwise 422 as a contact UUID.


@router.get("/companies", response_model=list[CompanyOut])
async def list_companies(current_user: CurrentUser, db: DB):
    return await service.list_companies(db, current_user.tenant_id)


@router.post("/companies", response_model=CompanyOut, status_code=status.HTTP_201_CREATED)
async def create_company(body: CompanyCreate, current_user: AdminUser, db: DB):
    try:
        return await service.create_company(db, current_user.tenant_id, body)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(e))


@router.get("/companies/{company_id}/contacts", response_model=list[CompanyContactOut])
async def list_company_contacts(company_id: uuid.UUID, current_user: CurrentUser, db: DB):
    company = await service.get_company(db, current_user.tenant_id, company_id)
    if not company:
        raise HTTPException(status_code=404, detail="Company not found")
    return await service.list_company_contacts(db, current_user.tenant_id, company_id)


@router.patch("/companies/{company_id}", response_model=CompanyOut)
async def update_company(company_id: uuid.UUID, body: CompanyUpdate, current_user: AdminUser, db: DB):
    company = await service.get_company(db, current_user.tenant_id, company_id)
    if not company:
        raise HTTPException(status_code=404, detail="Company not found")
    try:
        return await service.update_company(db, company, body)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(e))


@router.delete("/companies/{company_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_company(company_id: uuid.UUID, current_user: AdminUser, db: DB):
    company = await service.get_company(db, current_user.tenant_id, company_id)
    if not company:
        raise HTTPException(status_code=404, detail="Company not found")
    await service.delete_company(db, company)


@router.get("/labels", response_model=list[ContactLabelOut])
async def list_labels(current_user: CurrentUser, db: DB):
    return await service.list_labels(db, current_user.tenant_id)


@router.post("/labels", response_model=ContactLabelOut, status_code=status.HTTP_201_CREATED)
async def create_label(body: ContactLabelCreate, current_user: AdminUser, db: DB):
    try:
        return await service.create_label(db, current_user.tenant_id, body)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(e))


@router.patch("/labels/{label_id}", response_model=ContactLabelOut)
async def update_label(label_id: uuid.UUID, body: ContactLabelUpdate, current_user: AdminUser, db: DB):
    label = await service.get_label(db, current_user.tenant_id, label_id)
    if not label:
        raise HTTPException(status_code=404, detail="Label not found")
    try:
        return await service.update_label(db, label, body)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(e))


@router.delete("/labels/{label_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_label(label_id: uuid.UUID, current_user: AdminUser, db: DB):
    label = await service.get_label(db, current_user.tenant_id, label_id)
    if not label:
        raise HTTPException(status_code=404, detail="Label not found")
    await service.delete_label(db, label)


@router.post("/import/preview", response_model=ImportPreview)
async def import_preview(
    current_user: AdminUser,
    db: DB,
    file: UploadFile = File(...),
):
    content = await file.read()
    if len(content) > _MAX_IMPORT_SIZE:
        raise HTTPException(status_code=413, detail="File too large (max 10 MB)")
    mime = (file.content_type or "").split(";")[0].strip()
    if mime and mime not in _ALLOWED_IMPORT_TYPES:
        raise HTTPException(status_code=415, detail="Unsupported file type. Use CSV, JSON, or XLSX.")
    if not content:
        raise HTTPException(status_code=400, detail="Empty file")
    rows = _parse_import_file(file.filename or "", content)
    if not rows:
        raise HTTPException(status_code=400, detail="No rows found in file")
    # Headers in first-seen order across the parsed rows (rows may have ragged keys).
    headers: list[str] = []
    for row in rows:
        for k in row.keys():
            if k and k not in headers:
                headers.append(k)
    return ImportPreview(headers=headers, preview_rows=rows[:3])


_ALLOWED_IMPORT_TYPES = {
    "text/csv",
    "application/json",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_MAX_IMPORT_SIZE = 10_000_000  # 10 MB


@router.post("/import", response_model=ImportResult)
async def import_contacts(
    current_user: AdminUser,
    db: DB,
    file: UploadFile = File(...),
    column_mapping: Optional[str] = Form(None),
):
    content = await file.read()
    if len(content) > _MAX_IMPORT_SIZE:
        raise HTTPException(status_code=413, detail="File too large (max 10 MB)")
    mime = (file.content_type or "").split(";")[0].strip()
    if mime and mime not in _ALLOWED_IMPORT_TYPES:
        raise HTTPException(status_code=415, detail="Unsupported file type. Use CSV, JSON, or XLSX.")
    if not content:
        raise HTTPException(status_code=400, detail="Empty file")
    rows = _parse_import_file(file.filename or "", content)
    if not rows:
        raise HTTPException(status_code=400, detail="No rows found in file")

    if column_mapping:
        try:
            mapping = json.loads(column_mapping)
        except json.JSONDecodeError:
            raise HTTPException(status_code=400, detail="Invalid column_mapping JSON")
        if not isinstance(mapping, dict):
            raise HTTPException(status_code=400, detail="column_mapping must be an object")
        # incoming_col -> yippie_field; only known target fields are applied.
        pairs = [
            (str(src), dst)
            for src, dst in mapping.items()
            if dst in _IMPORT_TARGET_FIELDS
        ]
        rows = [
            {dst: row.get(src) for src, dst in pairs}
            for row in rows
        ]

    return await service.import_contacts(db, current_user.tenant_id, current_user.id, rows)


@router.get("/export")
async def export_contacts(
    current_user: CurrentUser,
    db: DB,
    search: Optional[str] = Query(None),
    ids: Optional[str] = Query(None, description="Comma-separated contact IDs to export"),
):
    contact_ids: Optional[list[uuid.UUID]] = None
    if ids:
        try:
            contact_ids = [uuid.UUID(x) for x in ids.split(",") if x.strip()]
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid contact id in 'ids'")
    csv_text = await service.export_contacts_csv(db, current_user.tenant_id, search, contact_ids)
    return StreamingResponse(
        iter([csv_text]),
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="contacts.csv"'},
    )


@router.get("/trash", response_model=list[ContactOut])
async def list_trash(current_user: AdminUser, db: DB):
    return await service.list_deleted_contacts(db, current_user.tenant_id)


@router.post("/{contact_id}/restore", response_model=ContactOut)
async def restore_contact(contact_id: uuid.UUID, current_user: AdminUser, db: DB):
    contact = await service.get_deleted_contact(db, current_user.tenant_id, contact_id)
    if not contact:
        raise HTTPException(status_code=404, detail="Contact not found in trash")
    await service.restore_contact(db, contact)
    return await service.get_contact(db, current_user.tenant_id, contact_id)


@router.delete("/{contact_id}/permanent", status_code=status.HTTP_204_NO_CONTENT)
async def permanently_delete_contact(contact_id: uuid.UUID, current_user: AdminUser, db: DB):
    contact = await service.get_deleted_contact(db, current_user.tenant_id, contact_id)
    if not contact:
        raise HTTPException(status_code=404, detail="Contact not found in trash")
    await service.permanently_delete_contact(db, contact)


@router.post("/{contact_id}/log-call/analyze", response_model=CallAnalyzeResponse)
async def analyze_call(contact_id: uuid.UUID, body: CallAnalyzeRequest, current_user: CurrentUser, db: DB):
    contact = await service.get_contact(db, current_user.tenant_id, contact_id)
    if not contact:
        raise HTTPException(status_code=404, detail="Contact not found")
    if body.outcome == "connected" and not body.transcript.strip():
        raise HTTPException(status_code=400, detail="Transcript is empty")
    return await service.analyze_call(contact, body)


@router.post("/{contact_id}/log-call", response_model=CallLogSaveResponse)
async def log_call(contact_id: uuid.UUID, body: CallLogSaveRequest, current_user: CurrentUser, db: DB):
    contact = await service.get_contact(db, current_user.tenant_id, contact_id)
    if not contact:
        raise HTTPException(status_code=404, detail="Contact not found")
    return await service.save_call_log(db, current_user.tenant_id, current_user, contact, body)


@router.get("/{contact_id}", response_model=ContactOut)
async def get_contact(contact_id: uuid.UUID, current_user: CurrentUser, db: DB):
    contact = await service.get_contact(db, current_user.tenant_id, contact_id)
    if not contact:
        raise HTTPException(status_code=404, detail="Contact not found")
    return contact


@router.patch("/{contact_id}", response_model=ContactOut)
async def update_contact(contact_id: uuid.UUID, body: ContactUpdate, current_user: CurrentUser, db: DB):
    contact = await service.get_contact(db, current_user.tenant_id, contact_id)
    if not contact:
        raise HTTPException(status_code=404, detail="Contact not found")
    return await service.update_contact(db, contact, body)


@router.delete("/{contact_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_contact(contact_id: uuid.UUID, current_user: CurrentUser, db: DB):
    contact = await service.get_contact(db, current_user.tenant_id, contact_id)
    if not contact:
        raise HTTPException(status_code=404, detail="Contact not found")
    await service.delete_contact(db, contact)
